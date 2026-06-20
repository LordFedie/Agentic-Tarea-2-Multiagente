from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as PdfImage,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "entrega"
ASSETS_DIR = OUTPUT_DIR / "assets"
DOCX_PATH = OUTPUT_DIR / "Informe_Tarea_2_Multiagente.docx"
PDF_PATH = OUTPUT_DIR / "Informe_Tarea_2_Multiagente.pdf"


def main():
    OUTPUT_DIR.mkdir(exist_ok=True)
    ASSETS_DIR.mkdir(exist_ok=True)

    diagram_path = ASSETS_DIR / "arquitectura_multiagente.png"
    build_architecture_diagram(diagram_path)

    document = Document()
    configure_document(document)
    build_cover(document)
    add_summary(document)
    add_architecture_section(document, diagram_path)
    add_stack_section(document)
    add_prompts_section(document)
    add_evaluation_section(document)
    add_conclusions_section(document)
    add_appendix_section(document)
    add_footer_page_numbers(document)
    document.save(DOCX_PATH)
    build_pdf_report(diagram_path)

    print(DOCX_PATH)
    print(PDF_PATH)


def configure_document(document: Document):
    section = document.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.bold = True

    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    styles["Heading 1"].paragraph_format.space_before = Pt(16)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)

    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    styles["Heading 2"].paragraph_format.space_before = Pt(12)
    styles["Heading 2"].paragraph_format.space_after = Pt(6)

    styles["Heading 3"].font.size = Pt(12)
    styles["Heading 3"].font.color.rgb = RGBColor(0x1F, 0x4D, 0x78)
    styles["Heading 3"].paragraph_format.space_before = Pt(8)
    styles["Heading 3"].paragraph_format.space_after = Pt(4)

    if "CaptionCustom" not in styles:
        caption_style = styles.add_style(
            "CaptionCustom",
            WD_STYLE_TYPE.PARAGRAPH
        )
        caption_style.font.name = "Calibri"
        caption_style.font.size = Pt(10)
        caption_style.font.italic = True
        caption_style.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        caption_style.paragraph_format.space_before = Pt(4)
        caption_style.paragraph_format.space_after = Pt(8)


def build_cover(document: Document):
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Informe del Proyecto\nAsistente Educativo Personal mediante un Sistema Multiagente")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Tarea 2 de Agentic AI\n"
        "Implementacion local en Python y Ollama"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(13)

    document.add_paragraph("")

    metadata = document.add_table(rows=4, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.style = "Table Grid"
    rows = [
        ("Proyecto", "Agentic-Tarea-2-Multiagente"),
        ("Modelo utilizado", "gemma3:4b"),
        ("Motor local", "Ollama"),
        ("Persistencia", "Archivo data/calendar.json"),
    ]
    for row_index, (label, value) in enumerate(rows):
        metadata.cell(row_index, 0).text = label
        metadata.cell(row_index, 1).text = value
        metadata.cell(row_index, 0).paragraphs[0].runs[0].bold = True

    document.add_paragraph("")
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(
        "Este informe documenta la arquitectura implementada, la ingenieria de prompts, "
        "el uso de herramientas locales, las pruebas realizadas durante desarrollo y "
        "las principales limitaciones observadas."
    )

    document.add_page_break()


def add_summary(document: Document):
    document.add_heading("1. Objetivo del proyecto", level=1)
    document.add_paragraph(
        "El objetivo del proyecto fue construir un asistente educativo personal basado en una "
        "arquitectura multiagente local. El sistema recibe consultas en lenguaje natural de "
        "estudiantes, las descompone cuando es necesario, delega la resolucion al agente "
        "especialista correspondiente y consolida una respuesta final con trazas visibles "
        "de ejecucion."
    )
    document.add_paragraph(
        "La solucion implementada incorpora cinco roles diferenciados: un agente coordinador "
        "como punto de entrada, un agente calculador para operaciones exactas, un agente "
        "organizador para la persistencia de eventos academicos, un agente experto para "
        "consultas conceptuales con apoyo de busqueda web y un agente critico para validar "
        "la respuesta final antes de entregarla al usuario."
    )
    document.add_paragraph(
        "Ademas de la funcionalidad basica exigida por la tarea, el sistema fue extendido "
        "durante el desarrollo para soportar consultas compuestas, normalizacion de fechas "
        "en espanol y reintentos automaticos cuando el agente critico detecta una respuesta "
        "invalida."
    )


def add_architecture_section(document: Document, diagram_path: Path):
    document.add_heading("2. Arquitectura multiagente", level=1)
    document.add_paragraph(
        "La arquitectura sigue un esquema jerarquico. El usuario interactua con un "
        "CoordinatorAgent que decide el flujo de trabajo. Este coordinador puede "
        "separar una consulta en varias subtareas, clasificar cada una por tipo y "
        "delegarla a uno de los agentes satelite. Una vez obtenida la respuesta, "
        "el CriticAgent valida la salida y puede solicitar un segundo intento con "
        "retroalimentacion explicita si detecta un problema."
    )

    document.add_picture(str(diagram_path), width=Inches(6.2))
    caption = document.add_paragraph(style="CaptionCustom")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.add_run(
        "Figura 1. Arquitectura general del sistema multiagente y relacion entre "
        "agentes y herramientas locales."
    )

    document.add_paragraph(
        "El CoordinatorAgent es responsable de la orquestacion. Primero recibe la "
        "consulta del usuario, luego detecta si la instruccion contiene una sola "
        "peticion o varias subtareas, y finalmente consulta al modelo para decidir "
        "si la resolucion debe pasar por el calculador, el organizador o el experto."
    )

    document.add_paragraph(
        "El CalculatorAgent opera sobre herramientas matematicas escritas en Python. "
        "Estas herramientas incluyen operaciones simples como suma, resta, multiplicacion, "
        "division, potencia y raiz cuadrada, asi como un evaluador seguro de expresiones "
        "aritmeticas completas basado en AST."
    )

    document.add_paragraph(
        "El OrganizerAgent utiliza un archivo JSON local como mecanismo de persistencia. "
        "Este agente puede crear, consultar, actualizar y eliminar eventos en "
        "data/calendar.json. Para mejorar la usabilidad, el sistema normaliza fechas "
        "escritas en espanol al formato ISO antes de invocar las herramientas de calendario."
    )

    document.add_paragraph(
        "El ExpertAgent resuelve preguntas conceptuales apoyandose en una busqueda "
        "web abierta con ddgs. El objetivo es reducir alucinaciones del modelo y "
        "forzar que la explicacion final se construya desde informacion recopilada "
        "por la herramienta."
    )

    document.add_paragraph(
        "El CriticAgent revisa la respuesta producida por cada subtarea. Si detecta "
        "una salida vacia, incoherente o que no cumple la accion solicitada, el "
        "CoordinatorAgent vuelve a invocar al especialista una vez mas, incorporando "
        "la retroalimentacion del critic en el segundo intento."
    )


def add_stack_section(document: Document):
    document.add_heading("3. Pila tecnologica y framework", level=1)
    document.add_paragraph(
        "La implementacion fue desarrollada completamente en Python, usando una "
        "arquitectura de agentes propia sobre Ollama. Esta decision se ajusta a la "
        "consigna, que permite libre eleccion del framework de agentes dentro del "
        "lenguaje Python, siempre que el sistema sea ejecutable localmente."
    )

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = table.rows[0].cells
    headers[0].text = "Componente"
    headers[1].text = "Herramienta"
    headers[2].text = "Justificacion"

    rows = [
        (
            "Modelo de lenguaje",
            "Ollama + gemma3:4b",
            "Permite ejecucion local con un modelo liviano y reproducible en hardware acotado."
        ),
        (
            "Orquestacion",
            "Clases Python",
            "Entrega control total del flujo entre agentes, herramientas y trazas de consola."
        ),
        (
            "Busqueda web",
            "ddgs",
            "Aporta una fuente abierta y gratuita para preguntas conceptuales del agente experto."
        ),
        (
            "Persistencia",
            "JSON local",
            "Cumple el requisito de almacenamiento persistente simple sin depender de una base externa."
        ),
        (
            "Evaluacion interna",
            "CriticAgent",
            "Introduce una capa de control de calidad y reintento antes de responder al usuario."
        ),
    ]

    for component, tool, reason in rows:
        row = table.add_row().cells
        row[0].text = component
        row[1].text = tool
        row[2].text = reason

    document.add_paragraph(
        "Aunque inicialmente se considero integrar Google ADK como framework adicional, "
        "se priorizo primero la resolucion correcta del flujo exigido por la tarea. "
        "La arquitectura actual funciona como una orquestacion multiagente propia y "
        "puede migrarse a una capa de framework mas formal como trabajo futuro, una vez "
        "estabilizados los prompts, las herramientas y la evaluacion del critic."
    )

    document.add_paragraph(
        "Las dependencias externas del proyecto quedaron reducidas a dos librerias "
        "principales: ollama para comunicarse con el modelo local y ddgs para la "
        "busqueda web del agente experto. Esto mantiene la instalacion simple y "
        "mejora la reproducibilidad del entorno."
    )


def add_prompts_section(document: Document):
    document.add_heading("4. Ingenieria de prompts", level=1)
    document.add_paragraph(
        "Cada agente dispone de un system prompt independiente. La estrategia general "
        "consiste en delimitar estrictamente el rol, las herramientas disponibles y "
        "el formato de salida esperado para reducir ambiguedad en modelos locales."
    )

    prompt_table = document.add_table(rows=1, cols=3)
    prompt_table.style = "Table Grid"
    prompt_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = prompt_table.rows[0].cells
    header[0].text = "Agente"
    header[1].text = "Responsabilidad del prompt"
    header[2].text = "Decision de diseno"

    entries = [
        (
            "CoordinatorAgent",
            "Clasificar la consulta y elegir el agente satelite adecuado.",
            "El prompt exige responder solo AGENT: Calculator, Organizer o Expert para simplificar el ruteo."
        ),
        (
            "CalculatorAgent",
            "Seleccionar una sola herramienta matematica o evaluar una expresion completa.",
            "Se prohibe calcular mentalmente para obligar el uso de funciones locales exactas."
        ),
        (
            "OrganizerAgent",
            "Traducir la consulta del usuario a una accion estructurada sobre el calendario.",
            "El formato TOOL/TITLE/DATE/TIME evita respuestas narrativas y facilita el parser."
        ),
        (
            "ExpertAgent",
            "Redactar explicaciones conceptuales usando informacion entregada por la busqueda.",
            "El prompt impide inventar fuentes o afirmar que se busco informacion inexistente."
        ),
        (
            "CriticAgent",
            "Validar si la respuesta es correcta y suficiente.",
            "Se enfatiza que no debe castigar diferencias solo de formato como 255 y 255.0."
        ),
    ]

    for agent, responsibility, decision in entries:
        row = prompt_table.add_row().cells
        row[0].text = agent
        row[1].text = responsibility
        row[2].text = decision

    document.add_paragraph(
        "La principal dificultad de la ingenieria de prompts estuvo en hacer que un "
        "modelo local pequeno respetara formatos cerrados con suficiente consistencia. "
        "Por esa razon, la implementacion combina prompts estrictos con parsers "
        "deterministas y validaciones adicionales en Python."
    )

    document.add_paragraph(
        "Un segundo aprendizaje importante fue que el critic no puede confiarse por "
        "completo al modelo. Durante el desarrollo aparecio un falso negativo donde "
        "la respuesta 255.0 fue rechazada a pesar de ser numericamente equivalente a 255. "
        "Esto obligo a reforzar el prompt del critic y agregar una correccion determinista "
        "para disputas de formato numerico."
    )


def add_evaluation_section(document: Document):
    document.add_heading("5. Evaluacion", level=1)
    document.add_paragraph(
        "La evaluacion se organizo en consultas exitosas y fallos observados durante "
        "el desarrollo. Dado que la solucion depende de un modelo local ejecutado con "
        "Ollama, para la entrega final se recomienda adjuntar las capturas o transcripciones "
        "exactas de consola obtenidas en la maquina del autor. En este informe se documenta "
        "el flujo esperado y los resultados observados durante la implementacion."
    )

    document.add_heading("5.1 Ejecuciones exitosas", level=2)

    successes = [
        (
            "Caso 1: Agente Calculador",
            "Consulta utilizada: Cual es el resultado de dividir 1024 entre 8?",
            [
                "El CoordinatorAgent clasifica la consulta como matematica y delega al CalculatorAgent.",
                "El CalculatorAgent selecciona la herramienta dividir.",
                "La operacion se resuelve con una funcion local exacta de Python y evita alucinaciones.",
                "El CriticAgent valida la consistencia del resultado final.",
            ],
            "Resultado esperado: respuesta final con resultado 128 o 128.0 y traza visible de la herramienta usada."
        ),
        (
            "Caso 2: Agente Organizador",
            "Consulta utilizada: Agrega una reunion de estudio de Fisica el 10 de junio de 2026 a las 15:00 horas.",
            [
                "El CoordinatorAgent detecta una tarea de calendario y la envia al OrganizerAgent.",
                "La fecha en espanol se normaliza a 2026-06-10 antes del llamado al modelo.",
                "El OrganizerAgent genera la accion add_event con titulo, fecha y hora.",
                "La herramienta local guarda el evento en data/calendar.json y el critic valida la salida.",
            ],
            "Resultado esperado: el evento queda persistido en el archivo JSON y aparece la traza de escritura en consola."
        ),
        (
            "Caso 3: Agente Experto",
            "Consulta utilizada: Que es la entropia en termodinamica?",
            [
                "El CoordinatorAgent clasifica la consulta como conceptual y activa al ExpertAgent.",
                "El ExpertAgent invoca la herramienta de busqueda web para recuperar contexto.",
                "Con los resultados obtenidos, el modelo redacta una explicacion academica breve y correcta.",
                "El CriticAgent revisa la respuesta y la marca como valida.",
            ],
            "Resultado esperado: respuesta conceptual clara, apoyada en busqueda web y con trazas de la herramienta."
        ),
    ]

    for title, query, steps, result in successes:
        document.add_heading(title, level=3)
        document.add_paragraph(query)
        for step in steps:
            document.add_paragraph(step, style="List Bullet")
        document.add_paragraph(result)

    document.add_heading("5.2 Ejecuciones deficientes o fallos", level=2)

    failures = [
        (
            "Fallo 1: Falso negativo del CriticAgent por formato numerico",
            "Consulta de desarrollo: Calcula el valor de 3 elevado a 5, mas 12.",
            "Durante una prueba de desarrollo, el sistema produjo correctamente 255.0, pero el critic rechazo la respuesta con el argumento de que el resultado correcto debia ser 255 y no 255.0.",
            "El problema se origino porque el critic estaba evaluando el formato textual del numero en lugar de su equivalencia semantica. Esto generaba reintentos innecesarios y podia degradar la experiencia de uso.",
            "Se corrigio reforzando el prompt del critic y agregando una validacion en Python que normaliza disputas de formato numerico cuando los valores son equivalentes."
        ),
        (
            "Fallo 2: Limitacion inicial ante consultas compuestas y expresiones de varios pasos",
            "Consulta de desarrollo: Calcula 25 x (4 + 6) y agrega una reunion de estudio de Fisica el 10 de junio de 2026 a las 15:00 horas.",
            "La primera version del sistema solo permitia enrutar la consulta completa a un unico agente. Esto impedia resolver instrucciones compuestas y tambien hacia fragil el manejo de expresiones matematicas con varias operaciones.",
            "La causa era una arquitectura de coordinacion demasiado simple: el coordinador clasificaba una sola intencion y el calculador estaba limitado a elegir una unica operacion atomica.",
            "La solucion fue introducir un descomponedor de consultas, un evaluador seguro de expresiones completas y un flujo de consolidacion de respuestas por subtarea."
        ),
    ]

    for title, query, observed, diagnosis, fix in failures:
        document.add_heading(title, level=3)
        document.add_paragraph(query)
        document.add_paragraph("Comportamiento observado: " + observed)
        document.add_paragraph("Diagnostico: " + diagnosis)
        document.add_paragraph("Correccion aplicada: " + fix)

    document.add_paragraph(
        "Como extensiones de evaluacion, tambien es recomendable conservar en el anexo "
        "las capturas de una consulta de lectura del calendario y una consulta compuesta "
        "que muestre explicitamente la deteccion de dos subtareas y el veredicto del critic "
        "para cada una."
    )


def add_conclusions_section(document: Document):
    document.add_heading("6. Conclusiones", level=1)
    document.add_paragraph(
        "El proyecto logro construir un sistema multiagente funcional y ejecutable localmente, "
        "capaz de combinar razonamiento basado en LLM con herramientas deterministas. La "
        "separacion de responsabilidades entre coordinador, especialistas y critic permite "
        "controlar mejor el flujo que una unica llamada directa al modelo."
    )

    document.add_paragraph(
        "Uno de los desafios principales fue la transferencia de contexto entre agentes. "
        "Los modelos locales pequenos tienden a desviarse del formato pedido o a simplificar "
        "demasiado la tarea. Por ello fue necesario complementar la ingenieria de prompts "
        "con parsers, reglas de validacion y funciones auxiliares en Python."
    )

    document.add_paragraph(
        "Otra limitacion importante fue la confiabilidad del critic. Aunque su presencia "
        "mejora el control de calidad, tambien introdujo errores de evaluacion cuando el "
        "modelo confundio equivalencia numerica con diferencia real. Esto evidencia que "
        "la supervision automatica de un sistema multiagente no puede depender solo de "
        "prompts y debe apoyarse en validaciones programaticas."
    )

    document.add_heading("7. Mejoras futuras", level=1)
    future_points = [
        "Integrar Google ADK u otro framework de agentes una vez estabilizada la logica base.",
        "Ampliar el parser de fechas para manejar expresiones relativas como manana o proximo lunes.",
        "Agregar pruebas automatizadas para parsers, calendario y evaluador matematico.",
        "Incorporar memoria de conversacion o contexto de usuario para consultas encadenadas.",
        "Mejorar la robustez del ExpertAgent mediante filtrado de fuentes y citas mas explicitas.",
    ]
    for point in future_points:
        document.add_paragraph(point, style="List Bullet")


def add_appendix_section(document: Document):
    document.add_heading("Anexo A. Consultas sugeridas para evidencias de consola", level=1)
    document.add_paragraph(
        "Las siguientes consultas son utiles para generar las capturas finales que se pueden "
        "adjuntar junto con este informe como evidencia de trazas de ejecucion:"
    )

    queries = [
        "Cual es el resultado de dividir 1024 entre 8?",
        "Agrega una reunion de estudio de Fisica el 10 de junio de 2026 a las 15:00 horas.",
        "Tengo algun evento programado para el 10 de junio de 2026?",
        "Que es la entropia en termodinamica?",
        "Calcula 25 x (4 + 6) y agrega una reunion de estudio de Fisica el 10 de junio de 2026 a las 15:00 horas.",
    ]
    for query in queries:
        document.add_paragraph(query, style="List Bullet")

    document.add_paragraph(
        "En cada captura conviene mostrar al menos: consulta de entrada, agente elegido, "
        "herramienta usada, veredicto del critic y respuesta final."
    )


def add_footer_page_numbers(document: Document):
    for section in document.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_page_number_field(paragraph)


def add_page_number_field(paragraph):
    run = paragraph.add_run("Pagina ")
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def build_architecture_diagram(output_path: Path):
    width, height = 1500, 900
    image = Image.new("RGB", (width, height), (248, 250, 252))
    draw = ImageDraw.Draw(image)

    try:
        title_font = ImageFont.truetype("arial.ttf", 34)
        box_font = ImageFont.truetype("arial.ttf", 26)
        small_font = ImageFont.truetype("arial.ttf", 22)
    except OSError:
        title_font = ImageFont.load_default()
        box_font = ImageFont.load_default()
        small_font = ImageFont.load_default()

    draw.text((420, 35), "Arquitectura del Sistema Multiagente", fill=(11, 37, 69), font=title_font)

    user_box = (590, 110, 910, 190)
    coordinator_box = (500, 250, 1000, 360)
    calc_box = (100, 450, 430, 590)
    org_box = (585, 450, 915, 590)
    expert_box = (1070, 450, 1400, 590)
    critic_box = (500, 665, 1000, 775)

    draw_round_box(draw, user_box, "Usuario", box_font, fill=(225, 239, 255))
    draw_round_box(draw, coordinator_box, "CoordinatorAgent\nClasificacion, descomposicion y consolidacion", box_font, fill=(203, 224, 247))
    draw_round_box(draw, calc_box, "CalculatorAgent\nHerramientas matematicas", box_font, fill=(228, 245, 231))
    draw_round_box(draw, org_box, "OrganizerAgent\nCalendario local JSON", box_font, fill=(255, 243, 205))
    draw_round_box(draw, expert_box, "ExpertAgent\nBusqueda web + explicacion", box_font, fill=(245, 232, 255))
    draw_round_box(draw, critic_box, "CriticAgent\nValidacion y reintento", box_font, fill=(255, 224, 224))

    draw_arrow(draw, (750, 190), (750, 250))
    draw_arrow(draw, (610, 360), (265, 450))
    draw_arrow(draw, (750, 360), (750, 450))
    draw_arrow(draw, (890, 360), (1235, 450))
    draw_arrow(draw, (265, 590), (610, 665))
    draw_arrow(draw, (750, 590), (750, 665))
    draw_arrow(draw, (1235, 590), (890, 665))
    draw_arrow(draw, (750, 775), (750, 840))

    draw.text((610, 820), "Respuesta final al usuario", fill=(11, 37, 69), font=box_font)

    add_tool_notes(draw, small_font)
    image.save(output_path)


def draw_round_box(draw, box, text, font, fill):
    draw.rounded_rectangle(box, radius=24, fill=fill, outline=(90, 110, 130), width=3)
    x1, y1, x2, y2 = box
    lines = text.split("\n")
    line_height = font.size + 6 if hasattr(font, "size") else 24
    total_height = len(lines) * line_height
    y = y1 + ((y2 - y1) - total_height) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        line_width = bbox[2] - bbox[0]
        x = x1 + ((x2 - x1) - line_width) / 2
        draw.text((x, y), line, fill=(20, 20, 20), font=font)
        y += line_height


def draw_arrow(draw, start, end):
    draw.line([start, end], fill=(90, 110, 130), width=4)
    ex, ey = end
    draw.polygon(
        [(ex, ey), (ex - 12, ey - 24), (ex + 12, ey - 24)],
        fill=(90, 110, 130)
    )


def add_tool_notes(draw, font):
    notes = [
        ((85, 605), "sumar, restar, multiplicar,\ndividir, potencia, raiz,\nevaluar_expresion"),
        ((570, 605), "add_event, update_event,\ndelete_event, get_events_by_date,\ndelete_all_events"),
        ((1045, 605), "web_search + redaccion\na partir de contexto"),
    ]
    for position, text in notes:
        draw.multiline_text(position, text, fill=(70, 70, 70), font=font, spacing=4)


def build_pdf_report(diagram_path: Path):
    styles = get_pdf_styles()
    story = []

    story.append(Paragraph(
        "Informe del Proyecto<br/>Asistente Educativo Personal mediante un Sistema Multiagente",
        styles["TitleCustom"]
    ))
    story.append(Spacer(1, 0.15 * inch))
    story.append(Paragraph(
        "Tarea 2 de Agentic AI<br/>Implementacion local en Python y Ollama",
        styles["SubtitleCustom"]
    ))
    story.append(Spacer(1, 0.25 * inch))

    metadata = Table(
        [
            ["Proyecto", "Agentic-Tarea-2-Multiagente"],
            ["Modelo utilizado", "gemma3:4b"],
            ["Motor local", "Ollama"],
            ["Persistencia", "Archivo data/calendar.json"],
        ],
        colWidths=[1.8 * inch, 4.3 * inch]
    )
    metadata.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.75, colors.HexColor("#AAB4BE")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF4FB")),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(metadata)
    story.append(Spacer(1, 0.2 * inch))
    story.append(Paragraph(
        "Este informe documenta la arquitectura implementada, la ingenieria de prompts, "
        "el uso de herramientas locales, las pruebas realizadas durante desarrollo y "
        "las principales limitaciones observadas.",
        styles["Body"]
    ))
    story.append(PageBreak())

    add_pdf_section(story, styles, "1. Objetivo del proyecto", [
        "El objetivo del proyecto fue construir un asistente educativo personal basado en una arquitectura multiagente local. "
        "El sistema recibe consultas en lenguaje natural de estudiantes, las descompone cuando es necesario, delega la resolucion "
        "al agente especialista correspondiente y consolida una respuesta final con trazas visibles de ejecucion.",
        "La solucion implementada incorpora cinco roles diferenciados: un agente coordinador como punto de entrada, un agente "
        "calculador para operaciones exactas, un agente organizador para la persistencia de eventos academicos, un agente experto "
        "para consultas conceptuales con apoyo de busqueda web y un agente critico para validar la respuesta final antes de entregarla al usuario.",
        "Ademas de la funcionalidad basica exigida por la tarea, el sistema fue extendido durante el desarrollo para soportar "
        "consultas compuestas, normalizacion de fechas en espanol y reintentos automaticos cuando el agente critico detecta una respuesta invalida.",
    ])

    story.append(Paragraph("2. Arquitectura multiagente", styles["Heading1Custom"]))
    story.append(Paragraph(
        "La arquitectura sigue un esquema jerarquico. El usuario interactua con un CoordinatorAgent que decide el flujo de trabajo. "
        "Este coordinador puede separar una consulta en varias subtareas, clasificar cada una por tipo y delegarla a uno de los agentes "
        "satelite. Una vez obtenida la respuesta, el CriticAgent valida la salida y puede solicitar un segundo intento con retroalimentacion explicita.",
        styles["Body"]
    ))
    story.append(Spacer(1, 0.1 * inch))
    story.append(PdfImage(str(diagram_path), width=6.1 * inch, height=3.66 * inch))
    story.append(Spacer(1, 0.08 * inch))
    story.append(Paragraph(
        "Figura 1. Arquitectura general del sistema multiagente y relacion entre agentes y herramientas locales.",
        styles["Caption"]
    ))
    for text in [
        "El CoordinatorAgent es responsable de la orquestacion. Primero recibe la consulta del usuario, luego detecta si la instruccion contiene una sola peticion o varias subtareas, y finalmente consulta al modelo para decidir si la resolucion debe pasar por el calculador, el organizador o el experto.",
        "El CalculatorAgent opera sobre herramientas matematicas escritas en Python. Estas herramientas incluyen operaciones simples como suma, resta, multiplicacion, division, potencia y raiz cuadrada, asi como un evaluador seguro de expresiones aritmeticas completas basado en AST.",
        "El OrganizerAgent utiliza un archivo JSON local como mecanismo de persistencia. Este agente puede crear, consultar, actualizar y eliminar eventos en data/calendar.json. Para mejorar la usabilidad, el sistema normaliza fechas escritas en espanol al formato ISO antes de invocar las herramientas de calendario.",
        "El ExpertAgent resuelve preguntas conceptuales apoyandose en una busqueda web abierta con ddgs. El objetivo es reducir alucinaciones del modelo y forzar que la explicacion final se construya desde informacion recopilada por la herramienta.",
        "El CriticAgent revisa la respuesta producida por cada subtarea. Si detecta una salida vacia, incoherente o que no cumple la accion solicitada, el CoordinatorAgent vuelve a invocar al especialista una vez mas, incorporando la retroalimentacion del critic en el segundo intento.",
    ]:
        story.append(Paragraph(text, styles["Body"]))

    story.append(Paragraph("3. Pila tecnologica y framework", styles["Heading1Custom"]))
    for text in [
        "La implementacion fue desarrollada completamente en Python, usando una arquitectura de agentes propia sobre Ollama. Esta decision se ajusta a la consigna, que permite libre eleccion del framework de agentes dentro del lenguaje Python, siempre que el sistema sea ejecutable localmente.",
        "Aunque inicialmente se considero integrar Google ADK como framework adicional, se priorizo primero la resolucion correcta del flujo exigido por la tarea. La arquitectura actual funciona como una orquestacion multiagente propia y puede migrarse a una capa de framework mas formal como trabajo futuro, una vez estabilizados los prompts, las herramientas y la evaluacion del critic.",
        "Las dependencias externas del proyecto quedaron reducidas a dos librerias principales: ollama para comunicarse con el modelo local y ddgs para la busqueda web del agente experto. Esto mantiene la instalacion simple y mejora la reproducibilidad del entorno.",
    ]:
        story.append(Paragraph(text, styles["Body"]))

    stack_table = Table(
        [
            ["Componente", "Herramienta", "Justificacion"],
            ["Modelo de lenguaje", "Ollama + gemma3:4b", "Ejecucion local con un modelo liviano y reproducible en hardware acotado."],
            ["Orquestacion", "Clases Python", "Control total del flujo entre agentes, herramientas y trazas de consola."],
            ["Busqueda web", "ddgs", "Fuente abierta y gratuita para preguntas conceptuales del agente experto."],
            ["Persistencia", "JSON local", "Almacenamiento persistente simple sin base de datos externa."],
            ["Evaluacion interna", "CriticAgent", "Control de calidad y reintento antes de responder al usuario."],
        ],
        colWidths=[1.35 * inch, 1.5 * inch, 3.45 * inch]
    )
    stack_table.setStyle(get_pdf_table_style())
    story.append(stack_table)

    story.append(Paragraph("4. Ingenieria de prompts", styles["Heading1Custom"]))
    story.append(Paragraph(
        "Cada agente dispone de un system prompt independiente. La estrategia general consiste en delimitar estrictamente el rol, las herramientas disponibles y el formato de salida esperado para reducir ambiguedad en modelos locales.",
        styles["Body"]
    ))
    prompt_table = Table(
        [
            ["Agente", "Responsabilidad del prompt", "Decision de diseno"],
            ["CoordinatorAgent", "Clasificar la consulta y elegir el agente satelite adecuado.", "Responde solo AGENT: Calculator, Organizer o Expert para simplificar el ruteo."],
            ["CalculatorAgent", "Seleccionar una sola herramienta matematica o evaluar una expresion completa.", "Se prohibe calcular mentalmente para obligar el uso de funciones locales exactas."],
            ["OrganizerAgent", "Traducir la consulta del usuario a una accion estructurada sobre el calendario.", "El formato TOOL/TITLE/DATE/TIME evita respuestas narrativas y facilita el parser."],
            ["ExpertAgent", "Redactar explicaciones conceptuales usando informacion entregada por la busqueda.", "Se impide inventar fuentes o afirmar que se busco informacion inexistente."],
            ["CriticAgent", "Validar si la respuesta es correcta y suficiente.", "No debe castigar diferencias solo de formato como 255 y 255.0."],
        ],
        colWidths=[1.25 * inch, 2.15 * inch, 3.0 * inch]
    )
    prompt_table.setStyle(get_pdf_table_style())
    story.append(prompt_table)
    story.append(Paragraph(
        "La principal dificultad de la ingenieria de prompts estuvo en hacer que un modelo local pequeno respetara formatos cerrados con suficiente consistencia. Por esa razon, la implementacion combina prompts estrictos con parsers deterministas y validaciones adicionales en Python.",
        styles["Body"]
    ))
    story.append(Paragraph(
        "Un segundo aprendizaje importante fue que el critic no puede confiarse por completo al modelo. Durante el desarrollo aparecio un falso negativo donde la respuesta 255.0 fue rechazada a pesar de ser numericamente equivalente a 255. Esto obligo a reforzar el prompt del critic y agregar una correccion determinista para disputas de formato numerico.",
        styles["Body"]
    ))

    story.append(Paragraph("5. Evaluacion", styles["Heading1Custom"]))
    story.append(Paragraph(
        "La evaluacion se organizo en consultas exitosas y fallos observados durante el desarrollo. Dado que la solucion depende de un modelo local ejecutado con Ollama, para la entrega final se recomienda adjuntar las capturas o transcripciones exactas de consola obtenidas en la maquina del autor. En este informe se documenta el flujo esperado y los resultados observados durante la implementacion.",
        styles["Body"]
    ))
    story.append(Paragraph("5.1 Ejecuciones exitosas", styles["Heading2Custom"]))

    success_cases = [
        (
            "Caso 1: Agente Calculador",
            "Consulta utilizada: Cual es el resultado de dividir 1024 entre 8?",
            [
                "El CoordinatorAgent clasifica la consulta como matematica y delega al CalculatorAgent.",
                "El CalculatorAgent selecciona la herramienta dividir.",
                "La operacion se resuelve con una funcion local exacta de Python y evita alucinaciones.",
                "El CriticAgent valida la consistencia del resultado final.",
            ],
            "Resultado esperado: respuesta final con resultado 128 o 128.0 y traza visible de la herramienta usada."
        ),
        (
            "Caso 2: Agente Organizador",
            "Consulta utilizada: Agrega una reunion de estudio de Fisica el 10 de junio de 2026 a las 15:00 horas.",
            [
                "El CoordinatorAgent detecta una tarea de calendario y la envia al OrganizerAgent.",
                "La fecha en espanol se normaliza a 2026-06-10 antes del llamado al modelo.",
                "El OrganizerAgent genera la accion add_event con titulo, fecha y hora.",
                "La herramienta local guarda el evento en data/calendar.json y el critic valida la salida.",
            ],
            "Resultado esperado: el evento queda persistido en el archivo JSON y aparece la traza de escritura en consola."
        ),
        (
            "Caso 3: Agente Experto",
            "Consulta utilizada: Que es la entropia en termodinamica?",
            [
                "El CoordinatorAgent clasifica la consulta como conceptual y activa al ExpertAgent.",
                "El ExpertAgent invoca la herramienta de busqueda web para recuperar contexto.",
                "Con los resultados obtenidos, el modelo redacta una explicacion academica breve y correcta.",
                "El CriticAgent revisa la respuesta y la marca como valida.",
            ],
            "Resultado esperado: respuesta conceptual clara, apoyada en busqueda web y con trazas de la herramienta."
        ),
    ]
    for title, query, steps, result in success_cases:
        story.append(Paragraph(title, styles["Heading3Custom"]))
        story.append(Paragraph(query, styles["Body"]))
        story.append(make_bullet_list(steps, styles))
        story.append(Paragraph(result, styles["Body"]))

    story.append(Paragraph("5.2 Ejecuciones deficientes o fallos", styles["Heading2Custom"]))
    failure_cases = [
        (
            "Fallo 1: Falso negativo del CriticAgent por formato numerico",
            "Consulta de desarrollo: Calcula el valor de 3 elevado a 5, mas 12.",
            "Durante una prueba de desarrollo, el sistema produjo correctamente 255.0, pero el critic rechazo la respuesta con el argumento de que el resultado correcto debia ser 255 y no 255.0.",
            "El problema se origino porque el critic estaba evaluando el formato textual del numero en lugar de su equivalencia semantica. Esto generaba reintentos innecesarios y podia degradar la experiencia de uso.",
            "Se corrigio reforzando el prompt del critic y agregando una validacion en Python que normaliza disputas de formato numerico cuando los valores son equivalentes."
        ),
        (
            "Fallo 2: Limitacion inicial ante consultas compuestas y expresiones de varios pasos",
            "Consulta de desarrollo: Calcula 25 x (4 + 6) y agrega una reunion de estudio de Fisica el 10 de junio de 2026 a las 15:00 horas.",
            "La primera version del sistema solo permitia enrutar la consulta completa a un unico agente. Esto impedia resolver instrucciones compuestas y tambien hacia fragil el manejo de expresiones matematicas con varias operaciones.",
            "La causa era una arquitectura de coordinacion demasiado simple: el coordinador clasificaba una sola intencion y el calculador estaba limitado a elegir una unica operacion atomica.",
            "La solucion fue introducir un descomponedor de consultas, un evaluador seguro de expresiones completas y un flujo de consolidacion de respuestas por subtarea."
        ),
    ]
    for title, query, observed, diagnosis, fix in failure_cases:
        story.append(Paragraph(title, styles["Heading3Custom"]))
        story.append(Paragraph(query, styles["Body"]))
        story.append(Paragraph("Comportamiento observado: " + observed, styles["Body"]))
        story.append(Paragraph("Diagnostico: " + diagnosis, styles["Body"]))
        story.append(Paragraph("Correccion aplicada: " + fix, styles["Body"]))

    story.append(Paragraph(
        "Como extensiones de evaluacion, tambien es recomendable conservar en el anexo las capturas de una consulta de lectura del calendario y una consulta compuesta que muestre explicitamente la deteccion de dos subtareas y el veredicto del critic para cada una.",
        styles["Body"]
    ))

    story.append(Paragraph("6. Conclusiones", styles["Heading1Custom"]))
    for text in [
        "El proyecto logro construir un sistema multiagente funcional y ejecutable localmente, capaz de combinar razonamiento basado en LLM con herramientas deterministas. La separacion de responsabilidades entre coordinador, especialistas y critic permite controlar mejor el flujo que una unica llamada directa al modelo.",
        "Uno de los desafios principales fue la transferencia de contexto entre agentes. Los modelos locales pequenos tienden a desviarse del formato pedido o a simplificar demasiado la tarea. Por ello fue necesario complementar la ingenieria de prompts con parsers, reglas de validacion y funciones auxiliares en Python.",
        "Otra limitacion importante fue la confiabilidad del critic. Aunque su presencia mejora el control de calidad, tambien introdujo errores de evaluacion cuando el modelo confundio equivalencia numerica con diferencia real. Esto evidencia que la supervision automatica de un sistema multiagente no puede depender solo de prompts y debe apoyarse en validaciones programaticas.",
    ]:
        story.append(Paragraph(text, styles["Body"]))

    story.append(Paragraph("7. Mejoras futuras", styles["Heading1Custom"]))
    future_points = [
        "Integrar Google ADK u otro framework de agentes una vez estabilizada la logica base.",
        "Ampliar el parser de fechas para manejar expresiones relativas como manana o proximo lunes.",
        "Agregar pruebas automatizadas para parsers, calendario y evaluador matematico.",
        "Incorporar memoria de conversacion o contexto de usuario para consultas encadenadas.",
        "Mejorar la robustez del ExpertAgent mediante filtrado de fuentes y citas mas explicitas.",
    ]
    story.append(make_bullet_list(future_points, styles))

    story.append(Paragraph("Anexo A. Consultas sugeridas para evidencias de consola", styles["Heading1Custom"]))
    story.append(Paragraph(
        "Las siguientes consultas son utiles para generar las capturas finales que se pueden adjuntar junto con este informe como evidencia de trazas de ejecucion:",
        styles["Body"]
    ))
    appendix_queries = [
        "Cual es el resultado de dividir 1024 entre 8?",
        "Agrega una reunion de estudio de Fisica el 10 de junio de 2026 a las 15:00 horas.",
        "Tengo algun evento programado para el 10 de junio de 2026?",
        "Que es la entropia en termodinamica?",
        "Calcula 25 x (4 + 6) y agrega una reunion de estudio de Fisica el 10 de junio de 2026 a las 15:00 horas.",
    ]
    story.append(make_bullet_list(appendix_queries, styles))
    story.append(Paragraph(
        "En cada captura conviene mostrar al menos: consulta de entrada, agente elegido, herramienta usada, veredicto del critic y respuesta final.",
        styles["Body"]
    ))

    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        leftMargin=1.0 * inch,
        rightMargin=1.0 * inch,
        topMargin=1.0 * inch,
        bottomMargin=0.8 * inch,
    )
    doc.build(story, onFirstPage=draw_pdf_footer, onLaterPages=draw_pdf_footer)


def add_pdf_section(story, styles, heading, paragraphs):
    story.append(Paragraph(heading, styles["Heading1Custom"]))
    for text in paragraphs:
        story.append(Paragraph(text, styles["Body"]))


def get_pdf_styles():
    styles = getSampleStyleSheet()

    styles.add(ParagraphStyle(
        name="TitleCustom",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#0B2545"),
        spaceAfter=8,
    ))
    styles.add(ParagraphStyle(
        name="SubtitleCustom",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=12,
        leading=16,
        alignment=TA_CENTER,
        spaceAfter=14,
    ))
    styles.add(ParagraphStyle(
        name="Heading1Custom",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#2E74B5"),
        spaceBefore=12,
        spaceAfter=8,
    ))
    styles.add(ParagraphStyle(
        name="Heading2Custom",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#2E74B5"),
        spaceBefore=10,
        spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        name="Heading3Custom",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=14,
        textColor=colors.HexColor("#1F4D78"),
        spaceBefore=8,
        spaceAfter=4,
    ))
    styles.add(ParagraphStyle(
        name="Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=14,
        alignment=TA_JUSTIFY,
        spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        name="Caption",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9.5,
        leading=11,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#555555"),
        spaceAfter=8,
    ))
    styles.add(ParagraphStyle(
        name="BulletBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=13,
        leftIndent=0,
        spaceAfter=2,
    ))
    return styles


def make_bullet_list(items, styles):
    return ListFlowable(
        [
            ListItem(Paragraph(item, styles["BulletBody"]), leftIndent=10)
            for item in items
        ],
        bulletType="bullet",
        start="circle",
        leftIndent=20,
        spaceAfter=6,
    )


def get_pdf_table_style():
    return TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#AAB4BE")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF1F8")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("LEADING", (0, 0), (-1, -1), 11),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("SPACEAFTER", (0, 0), (-1, -1), 8),
    ])


def draw_pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 9)
    canvas.setFillColor(colors.HexColor("#555555"))
    canvas.drawRightString(
        doc.pagesize[0] - doc.rightMargin,
        0.5 * inch,
        f"Pagina {canvas.getPageNumber()}"
    )
    canvas.restoreState()


if __name__ == "__main__":
    main()
